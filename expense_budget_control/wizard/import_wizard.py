# -*- coding: utf-8 -*-
import base64
import csv
import io
import logging

from openpyxl import load_workbook

from odoo import api, fields, models, _
from odoo.exceptions import AccessError, UserError

from ..models.expense import as_text, get_value, parse_amount

_logger = logging.getLogger(__name__)


class ExpenseBudgetImportWizard(models.TransientModel):
    _name = 'expense.budget.import.wizard'
    _description = 'Importar catálogo, proyecciones o gastos desde Excel/CSV'

    import_type = fields.Selection([
        ('accounts', 'Cuentas y rubros'),
        ('projections', 'Proyecciones mensuales'),
        ('expenses', 'Gastos reales'),
    ], string='Tipo de importación', required=True,
       default=lambda self: ('accounts' if self.env.su or self.env.user.has_group(
           'expense_budget_control.group_expense_budget_manager') else 'projections'))
    company_id = fields.Many2one('res.company', string='Empresa', required=True,
                                 default=lambda self: self.env.company)
    budget_id = fields.Many2one('expense.budget', string='Presupuesto anual',
                                domain="[('company_id', '=', company_id), ('state', '=', 'draft')]")
    upload = fields.Binary(string='Archivo XLSX o CSV', attachment=False, required=True)
    filename = fields.Char(string='Nombre del archivo')

    @api.model
    def _read_rows(self, content, filename, import_type):
        name = (filename or '').lower()
        if name.endswith('.xlsx'):
            try:
                workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            except Exception as exc:
                raise UserError(_('El Excel no se pudo abrir: %s') % exc) from exc
            try:
                target = {'accounts': 'Cuentas', 'projections': 'Proyecciones',
                          'expenses': 'Gastos'}[import_type]
                worksheet = workbook[target] if target in workbook.sheetnames else workbook.active
                iterator = worksheet.iter_rows(values_only=True)
                headers = next(iterator, ())
                if not headers:
                    raise UserError(_('El archivo no tiene encabezados.'))
                rows = []
                for line_no, values in enumerate(iterator, start=2):
                    if not any(value is not None and str(value).strip() for value in values):
                        continue
                    if line_no > 20001:
                        raise UserError(_('Máximo de 20.000 filas por importación.'))
                    rows.append((line_no, dict(zip((as_text(h) for h in headers), values))))
                return rows
            finally:
                workbook.close()
        if name.endswith('.csv'):
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError as exc:
                raise UserError(_('El archivo CSV debe estar codificado en UTF-8.')) from exc
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
                reader = csv.DictReader(io.StringIO(text), dialect=dialect)
            except csv.Error as exc:
                raise UserError(_('Error leyendo encabezados CSV: %s') % exc) from exc
            rows = []
            for index, row in enumerate(reader, start=2):
                if index > 20001:
                    raise UserError(_('Máximo de 20.000 filas por importación.'))
                if any(value and str(value).strip() for value in row.values()):
                    rows.append((index, row))
            return rows
        raise UserError(_('Solo se aceptan archivos .xlsx o .csv.'))

    def action_import(self):
        self.ensure_one()
        is_manager = self.env.su or self.env.user.has_group(
            'expense_budget_control.group_expense_budget_manager')
        is_registrar = self.env.user.has_group(
            'expense_budget_control.group_expense_budget_registrar')
        if not is_manager and not (is_registrar and self.import_type == 'projections'):
            raise AccessError(_(
                'Solo el administrador puede importar cuentas y gastos. '
                'El usuario de registro únicamente puede importar proyecciones mensuales.'))
        if self.company_id not in self.env.user.company_ids:
            raise AccessError(_('No tiene acceso a la empresa elegida.'))
        if self.import_type == 'projections':
            if not self.budget_id or self.budget_id.company_id != self.company_id:
                raise UserError(_('Seleccione un presupuesto de la empresa.'))
            if self.budget_id.state != 'draft':
                raise UserError(_('Solo se pueden importar proyecciones en borrador.'))
        if not self.filename or not self.upload:
            raise UserError(_('Adjunte un archivo con extensión .xlsx o .csv.'))
        raw = base64.b64decode(self.upload)
        if len(raw) > 8 * 1024 * 1024:
            raise UserError(_('El archivo supera 8 MB.'))
        rows = self._read_rows(raw, self.filename, self.import_type)
        if not rows:
            raise UserError(_('El archivo no contiene filas para procesar.'))

        created = updated = unchanged = 0
        for line_number, item in rows:
            try:
                if self.import_type == 'accounts':
                    outcome = self._import_account(item)
                elif self.import_type == 'projections':
                    outcome = self._import_projection(item)
                else:
                    outcome = self._import_expense(item)
                if outcome == 'created':
                    created += 1
                elif outcome == 'updated':
                    updated += 1
                else:
                    unchanged += 1
            except (ValueError, UserError) as exc:
                # Al elevar error se revierte la transacción completa: importación atómica.
                raise UserError(_('Error en la fila %(line)s: %(error)s') % {
                    'line': line_number, 'error': exc,
                }) from exc

        return {'type': 'ir.actions.client', 'tag': 'display_notification', 'params': {
            'title': _('Importación completada'),
            'message': _('Insertados: %(a)s | Actualizados: %(b)s | Sin cambios: %(c)s') % {
                'a': created, 'b': updated, 'c': unchanged,
            },
            'type': 'success', 'sticky': False,
        }}

    def _import_account(self, item):
        category_code = as_text(get_value(item, 'codigo_rubro', 'rubro_codigo', required=True))
        category_name = as_text(get_value(item, 'nombre_rubro', 'rubro', required=True))
        account_code = as_text(get_value(item, 'codigo_cuenta', 'cuenta_codigo', required=True))
        account_name = as_text(get_value(item, 'nombre_cuenta', 'cuenta_nombre', required=True))
        Category = self.env['expense.budget.category'].with_company(self.company_id)
        category = Category.search([('company_id', '=', self.company_id.id),
                                    ('code', '=', category_code)], limit=1)
        if not category:
            category = Category.create({'code': category_code, 'name': category_name,
                                        'company_id': self.company_id.id})
        elif category.name != category_name:
            category.write({'name': category_name})
        Account = self.env['expense.budget.account'].with_company(self.company_id)
        account = Account.search([('company_id', '=', self.company_id.id),
                                  ('code', '=', account_code)], limit=1)
        values = {'name': account_name, 'category_id': category.id}
        if account:
            if account.name != account_name or account.category_id != category:
                account.write(values)
                return 'updated'
            return 'unchanged'
        Account.create(dict(values, code=account_code, company_id=self.company_id.id))
        return 'created'

    def _import_projection(self, item):
        code = as_text(get_value(item, 'codigo_cuenta', 'cuenta', required=True))
        raw_month = as_text(get_value(item, 'mes', 'month', required=True))
        names = {'enero': '1', 'febrero': '2', 'marzo': '3', 'abril': '4', 'mayo': '5',
                 'junio': '6', 'julio': '7', 'agosto': '8', 'septiembre': '9',
                 'octubre': '10', 'noviembre': '11', 'diciembre': '12'}
        month = names.get(raw_month.lower(), raw_month)
        try:
            month = str(int(month))
        except ValueError as exc:
            raise ValueError('Mes inválido: %s.' % raw_month) from exc
        if month not in [str(i) for i in range(1, 13)]:
            raise ValueError('El mes debe estar entre 1 y 12.')
        amount = parse_amount(get_value(item, 'proyectado', 'presupuesto', 'monto', required=True))
        if amount < 0:
            raise ValueError('El presupuesto debe ser mayor o igual a cero.')
        account = self.env['expense.budget.account'].search([
            ('company_id', '=', self.company_id.id), ('code', '=', code),
        ], limit=1)
        if not account:
            raise ValueError('La cuenta %s no existe: importe antes el catálogo.' % code)
        Line = self.env['expense.budget.line']
        line = Line.search([
            ('budget_id', '=', self.budget_id.id), ('month', '=', month),
            ('account_id', '=', account.id),
        ], limit=1)
        if line:
            if line.planned_amount != amount:
                line.write({'planned_amount': amount})
                return 'updated'
            return 'unchanged'
        Line.create({'budget_id': self.budget_id.id, 'account_id': account.id,
                     'month': month, 'planned_amount': amount})
        return 'created'

    def _import_expense(self, item):
        external_id = get_value(item, 'identificador', 'id', 'external_id', required=True)
        return self.env['expense.budget.expense'].with_company(self.company_id)._upsert_external(
            self.company_id, item, 'excel', external_id,
        )
