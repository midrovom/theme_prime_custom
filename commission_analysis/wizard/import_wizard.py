import base64
import hashlib
import io
import struct
import unicodedata
from datetime import date, datetime

from odoo import fields, models, _
from odoo.exceptions import UserError


MONTHS = {
    'ene': 1, 'feb': 2, 'mar': 3, 'abr': 4, 'may': 5, 'jun': 6,
    'jul': 7, 'ago': 8, 'sept': 9, 'sep': 9, 'oct': 10, 'nov': 11, 'dic': 12,
}
CHUNK = 2000


def normalize_text(value):
    text = str(value or '').strip()
    text = unicodedata.normalize('NFKD', text)
    return ''.join(ch for ch in text if not unicodedata.combining(ch)).casefold()


def make_icc_key(value, length=18):
    digits = ''.join(ch for ch in str(value or '').strip() if ch.isdigit())
    return digits[:length] if len(digits) >= length else digits


def normalize_percentage(value):
    """Return (source_value, percent_points). Claro exports 20% as numeric 0.20."""
    if value in (None, ''):
        return 0.0, 0.0
    try:
        if isinstance(value, str):
            clean = value.strip().replace('%', '').replace(',', '.')
            numeric = float(clean or 0)
            if '%' in value:
                return numeric / 100.0, numeric
        else:
            numeric = float(value)
        return numeric, numeric * 100.0 if abs(numeric) <= 1.0 and numeric != 0 else numeric
    except (TypeError, ValueError):
        return 0.0, 0.0


def parse_date(value):
    if not value:
        return False
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().lower().replace(',', '')
    for key, month in MONTHS.items():
        if text.startswith(key + ' '):
            parts = text.split()
            if len(parts) >= 3:
                try:
                    return datetime(int(parts[-1]), month, int(parts[-2])).date()
                except (ValueError, TypeError):
                    return False
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y'):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return False


def dbf_rows(data):
    """Minimal dBase/FoxBase reader for the supplier purchase export."""
    if len(data) < 64:
        raise UserError(_('El archivo de compras no tiene una estructura DBF válida.'))
    header_len = struct.unpack('<H', data[8:10])[0]
    record_len = struct.unpack('<H', data[10:12])[0]
    record_count = struct.unpack('<I', data[4:8])[0]
    if header_len <= 32 or record_len <= 1 or header_len > len(data):
        raise UserError(_('Cabecera DBF inválida en el archivo de compras.'))

    field_specs = []
    pos = 32
    while pos + 32 <= header_len and data[pos] != 0x0D:
        descriptor = data[pos:pos + 32]
        field_name = descriptor[:11].split(b'\x00')[0].decode('latin1').strip()
        field_type = chr(descriptor[11])
        field_len = descriptor[16]
        decimals = descriptor[17]
        field_specs.append((field_name, field_type, field_len, decimals))
        pos += 32

    max_records = max(0, (len(data) - header_len) // record_len)
    record_count = min(record_count, max_records)
    for idx in range(record_count):
        record = data[header_len + idx * record_len:header_len + (idx + 1) * record_len]
        if not record or record[:1] == b'*':
            continue
        values = {}
        cursor = 1
        for field_name, field_type, field_len, _decimals in field_specs:
            raw = record[cursor:cursor + field_len]
            cursor += field_len
            text = raw.decode('latin1', 'ignore').strip()
            if field_type in 'NF':
                try:
                    values[field_name] = float(text) if text else 0.0
                except ValueError:
                    values[field_name] = 0.0
            elif field_type == 'D' and len(text) == 8:
                try:
                    values[field_name] = datetime.strptime(text, '%Y%m%d').date()
                except ValueError:
                    values[field_name] = False
            else:
                values[field_name] = text
        yield idx + 1, values


class CommissionImportWizard(models.TransientModel):
    _name = 'commission.import.wizard'
    _description = 'Importar Compras / Liquidaciones'

    import_type = fields.Selection(
        [('purchase', 'Compras'), ('settlement', 'Liquidación Claro')],
        required=True, default='settlement',
    )
    purchase_batch_id = fields.Many2one('commission.purchase.batch')
    settlement_batch_id = fields.Many2one('commission.settlement.batch')
    file_ids = fields.Many2many('ir.attachment', string='Archivos', required=True)

    def action_import(self):
        self.ensure_one()
        if self.import_type == 'purchase':
            if not self.purchase_batch_id:
                raise UserError(_('Seleccione un lote de compras.'))
            if self.purchase_batch_id.state == 'closed':
                raise UserError(_('No se puede importar en un lote de compras cerrado.'))
        else:
            if not self.settlement_batch_id:
                raise UserError(_('Seleccione una liquidación.'))
            if self.settlement_batch_id.state == 'closed':
                raise UserError(_('No se puede importar en una liquidación cerrada.'))

        ImportFile = self.env['commission.import.file']
        for attachment in self.file_ids:
            if not attachment.datas:
                raise UserError(_('El archivo %s no contiene datos.') % attachment.name)
            data = base64.b64decode(attachment.datas)
            sha256 = hashlib.sha256(data).hexdigest()
            if ImportFile.search_count([('sha256', '=', sha256)]):
                raise UserError(_('El archivo %s ya fue importado anteriormente.') % attachment.name)

            if self.import_type == 'purchase':
                count = self._import_purchase(attachment.name, data)
            else:
                count = self._import_settlement(attachment.name, data)

            ImportFile.create({
                'name': attachment.name,
                'sha256': sha256,
                'attachment_id': attachment.id,
                'purchase_batch_id': self.purchase_batch_id.id if self.import_type == 'purchase' else False,
                'settlement_batch_id': self.settlement_batch_id.id if self.import_type == 'settlement' else False,
                'row_count': count,
            })
        return {'type': 'ir.actions.client', 'tag': 'reload'}

    def _warehouse_mapping(self, code, name):
        Mapping = self.env['commission.warehouse.mapping']
        code = str(code or '').strip()
        name = str(name or '').strip()
        domain = [('active', '=', True)]
        if code and name:
            domain += ['|', ('warehouse_code', '=ilike', code), ('warehouse_name', '=ilike', name)]
        elif code:
            domain += [('warehouse_code', '=ilike', code)]
        elif name:
            domain += [('warehouse_name', '=ilike', name)]
        else:
            return Mapping.browse()
        return Mapping.search(domain, limit=1)

    def _import_purchase(self, name, data):
        batch = self.purchase_batch_id
        operator = batch.operator_id
        key_length = operator.icc_match_length or 18
        Sim = self.env['commission.sim']
        Line = self.env['commission.purchase.line']
        count = 0
        buffer = []
        mapping_cache = {}
        affected_sim_ids = set()

        def get_mapping(code, warehouse_name):
            key = (normalize_text(code), normalize_text(warehouse_name))
            if key not in mapping_cache:
                mapping_cache[key] = self._warehouse_mapping(code, warehouse_name)
            return mapping_cache[key]

        def flush(rows):
            nonlocal count
            if not rows:
                return
            keys = list({vals['icc_key'] for vals in rows if vals['icc_key']})
            existing = {
                sim.icc_key: sim
                for sim in Sim.search([
                    ('operator_id', '=', operator.id),
                    ('icc_key', 'in', keys),
                ])
            }
            new_values = []
            seen_new = set()
            for vals in rows:
                key = vals['icc_key']
                if not key:
                    continue
                if key not in existing and key not in seen_new:
                    seen_new.add(key)
                    mapping = get_mapping(vals['warehouse_code'], vals['warehouse_name'])
                    new_values.append({
                        'icc': vals['icc'],
                        'icc_key': key,
                        'operator_id': operator.id,
                        'purchase_date': vals['date'],
                        'invoice': vals['invoice'],
                        'warehouse': vals['warehouse_name'],
                        'product': vals['product_name'],
                        'purchase_cost': operator._roi_cost_value(
                            vals['cost'], vals['discount'], vals['tax']
                        ),
                        'region_id': mapping.region_id.id if mapping else False,
                        'zone_id': mapping.zone_id.id if mapping else False,
                    })
            if new_values:
                for sim in Sim.create(new_values):
                    existing[sim.icc_key] = sim

            line_values = []
            for vals in rows:
                sim = existing.get(vals['icc_key'])
                if not sim:
                    continue
                mapping = get_mapping(vals['warehouse_code'], vals['warehouse_name'])
                vals.update({
                    'sim_id': sim.id,
                    'region_id': mapping.region_id.id if mapping else sim.region_id.id,
                    'zone_id': mapping.zone_id.id if mapping else sim.zone_id.id,
                })
                line_values.append(vals)
                affected_sim_ids.add(sim.id)
            if line_values:
                Line.create(line_values)
                count += len(line_values)

        for row_number, row in dbf_rows(data):
            icc = str(row.get('CHIP') or '').strip()
            key = make_icc_key(icc, key_length)
            if not icc or not key:
                continue
            buffer.append({
                'icc': icc,
                'icc_key': key,
                'date': row.get('FECHA'),
                'document_type': row.get('TIPO'),
                'number': row.get('NUMERO'),
                'supplier_code': row.get('PROVEEDOR'),
                'supplier_name': row.get('NPROVEEDOR'),
                'invoice': row.get('FACTURA'),
                'warehouse_code': row.get('BODEGA'),
                'warehouse_name': row.get('NBODEGA'),
                'product_code': row.get('PRODUCTO'),
                'product_name': row.get('NPRODUCTO'),
                'quantity': row.get('CANTIDAD') or 1,
                'cost': row.get('COSTO') or 0,
                'discount': row.get('DESCUENTO') or 0,
                'tax': row.get('IMPUESTO') or 0,
                'source_file': name,
                'source_row': row_number,
                'batch_id': batch.id,
            })
            if len(buffer) >= CHUNK:
                flush(buffer)
                buffer = []
        flush(buffer)

        sim_ids = list(affected_sim_ids)
        for start in range(0, len(sim_ids), 10000):
            chunk = sim_ids[start:start + 10000]
            batch._recalculate_roi_cost_for_sims(chunk)
            Sim._refresh_kpis(chunk)
        batch.state = 'processed'
        return count

    @staticmethod
    def _build_header_map(top, second):
        # Claro uses two header rows: descriptive columns in row 2 and Measures in row 1.
        size = max(len(top), len(second))
        headers = []
        for idx in range(size):
            second_value = second[idx] if idx < len(second) else None
            top_value = top[idx] if idx < len(top) else None
            headers.append(second_value or top_value)
        return {normalize_text(value): idx for idx, value in enumerate(headers) if value not in (None, '')}

    def _import_settlement(self, name, data):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise UserError(_('Se requiere la librería Python openpyxl.')) from exc

        batch = self.settlement_batch_id
        operator = batch.operator_id
        key_length = operator.icc_match_length or 18
        try:
            worksheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        except Exception as exc:
            raise UserError(_('No se pudo leer %s como XLSX: %s') % (name, exc)) from exc

        top = list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        second = list(next(worksheet.iter_rows(min_row=2, max_row=2, values_only=True)))
        header_map = self._build_header_map(top, second)

        aliases = {
            'compensation_date': ['compensation date'],
            'region': ['region', 'región'],
            'customer': ['cliente'],
            'product': ['producto'],
            'concept': ['concepto'],
            'service_number': ['numero servicio', 'número servicio'],
            'activation_date': ['fecha activacion', 'fecha activación'],
            'regularization_date': ['fecha regularizacion', 'fecha regularización'],
            'high_date': ['fecha alta'],
            'distributor_type': ['tipo distribuidor'],
            'invoice': ['no. factura', 'no factura'],
            'simcard': ['simcard'],
            'pvp': ['pvp'],
            'percentage': ['porcentaje'],
            'recharges': ['recargas'],
            'consumptions': ['consumos'],
            'base_tariff': ['tarifa basica', 'tarifa básica'],
            'evaluated_consumption': ['consumo evaluado'],
            'value': ['valor'],
            'discounted_base_tariff': ['tarifa basica descontada', 'tarifa básica descontada'],
            'advance_value': ['valor anticipo'],
        }

        indexes = {}
        for logical, names in aliases.items():
            found = None
            for alias in names:
                key = normalize_text(alias)
                if key in header_map:
                    found = header_map[key]
                    break
            indexes[logical] = found

        required = ['compensation_date', 'region', 'concept', 'simcard', 'percentage', 'evaluated_consumption', 'value']
        missing = [field for field in required if indexes.get(field) is None]
        if missing:
            raise UserError(_(
                'Faltan columnas requeridas en %s: %s'
            ) % (name, ', '.join(missing)))

        Line = self.env['commission.settlement.line']
        Sim = self.env['commission.sim']
        Region = self.env['commission.region']
        count = 0
        buffer = []
        region_cache = {}
        affected_sim_ids = set()
        rules = Line._rules_for_operator(operator.id)

        def get(row, logical):
            idx = indexes.get(logical)
            return row[idx] if idx is not None and idx < len(row) else None

        def get_region(region_name):
            key = str(region_name or '').strip()
            if not key:
                return Region.browse()
            cache_key = normalize_text(key)
            if cache_key not in region_cache:
                region_cache[cache_key] = (
                    Region.search([('name', '=ilike', key)], limit=1)
                    or Region.create({'name': key})
                )
            return region_cache[cache_key]

        def flush(rows):
            nonlocal count
            if not rows:
                return
            keys = list({vals['icc_key'] for vals in rows if vals['icc_key']})
            sims = {
                sim.icc_key: sim
                for sim in Sim.search([
                    ('operator_id', '=', operator.id),
                    ('icc_key', 'in', keys),
                ])
            }
            fallback_region_updates = {}
            for vals in rows:
                sim = sims.get(vals['icc_key'])
                if sim:
                    vals['sim_id'] = sim.id
                    affected_sim_ids.add(sim.id)
                    # Telecity's analytical region/zone from purchase mapping has priority.
                    source_region_id = vals.get('region_id')
                    vals['region_id'] = sim.region_id.id or source_region_id
                    vals['zone_id'] = sim.zone_id.id
                    # If the purchase master has no region yet, use Claro's source Region as a
                    # fallback. A later Bodega -> Región/Zona mapping is allowed to overwrite it.
                    if not sim.region_id and source_region_id:
                        fallback_region_updates.setdefault(source_region_id, []).append(sim.id)
                vals.update(Line._expected_values_from_dict(vals, operator.id, rules=rules))
            Line.create(rows)
            for region_id, sim_ids in fallback_region_updates.items():
                self.env.cr.execute(
                    "UPDATE commission_sim SET region_id=%s WHERE id = ANY(%s) AND region_id IS NULL",
                    (region_id, list(set(sim_ids))),
                )
            if fallback_region_updates:
                Sim.invalidate_model(['region_id'])
            count += len(rows)

        for row_number, row in enumerate(worksheet.iter_rows(min_row=3, values_only=True), 3):
            icc = str(get(row, 'simcard') or '').strip()
            key = make_icc_key(icc, key_length)
            if not icc or not key:
                continue
            source_percentage, normalized_percentage = normalize_percentage(get(row, 'percentage'))
            source_region = get(row, 'region')
            region = get_region(source_region)
            buffer.append({
                'batch_id': batch.id,
                'icc': icc,
                'icc_key': key,
                'compensation_date': parse_date(get(row, 'compensation_date')),
                'region_text': source_region,
                'region_id': region.id if region else False,
                'customer': get(row, 'customer'),
                'product': get(row, 'product'),
                'concept': get(row, 'concept'),
                'service_number': get(row, 'service_number'),
                'activation_date': parse_date(get(row, 'activation_date')),
                'regularization_date': parse_date(get(row, 'regularization_date')),
                'high_date': parse_date(get(row, 'high_date')),
                'distributor_type': get(row, 'distributor_type'),
                'invoice': get(row, 'invoice'),
                'pvp': get(row, 'pvp') or 0,
                'reported_percentage_source': source_percentage,
                'reported_percentage': normalized_percentage,
                'recharges': get(row, 'recharges') or 0,
                'consumptions': get(row, 'consumptions') or 0,
                'base_tariff': get(row, 'base_tariff') or 0,
                'evaluated_consumption': get(row, 'evaluated_consumption') or 0,
                'reported_value': get(row, 'value') or 0,
                'discounted_base_tariff': get(row, 'discounted_base_tariff') or 0,
                'advance_value': get(row, 'advance_value') or 0,
                'source_file': name,
                'source_row': row_number,
            })
            if len(buffer) >= CHUNK:
                flush(buffer)
                buffer = []
        flush(buffer)

        sim_ids = list(affected_sim_ids)
        for start in range(0, len(sim_ids), 10000):
            Sim._refresh_kpis(sim_ids[start:start + 10000])
        batch.state = 'processed'
        return count
